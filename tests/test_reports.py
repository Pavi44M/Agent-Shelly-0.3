"""v1.1: real-data intake, report engine, Excel workbooks (dashboard, formulas, pivots)."""
import warnings
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
import yaml

from shelly import ingest

ROOT = Path(__file__).resolve().parent.parent
CFG = yaml.safe_load(open(ROOT / "config.yaml"))


@pytest.fixture(scope="module")
def ctx(tmp_path_factory):
    import contextlib
    import io
    from shelly.agent import run
    from shelly.reports.context import Ctx
    out = tmp_path_factory.mktemp("run")
    with contextlib.redirect_stdout(io.StringIO()):
        res = run(str(ROOT / "data/sample"), str(out))
    return Ctx(res["prepared"], res["results"], CFG, res["actions"], res["summary"], "test", "Test store")


def test_messy_export_is_matched(tmp_path):
    s = pd.read_csv(ROOT / "data/sample/sales.csv").head(3000).merge(pd.read_csv(ROOT / "data/sample/products.csv"), on="sku")
    pd.DataFrame({"Trans Date": pd.to_datetime(s["date"]).dt.strftime("%d/%m/%Y"), "Article No": s["sku"], "Description": s["product_name"],
                  "Department": s["category"], "Qty Sold": s["units"], "Net Amt excl GST": s["net_sales"].map(lambda v: f"${v:,.2f}")}
                 ).to_csv(tmp_path / "POS export.csv", index=False, sep=";")
    tables, maps = ingest.load_folder(tmp_path, {}, {"default_margin_pct": 30})
    m = {k: v[0] for k, v in maps[0].columns.items()}
    assert m["date"] == "Trans Date" and m["sku"] == "Article No" and m["units"] == "Qty Sold" and m["net_sales"] == "Net Amt excl GST"
    assert "products" in tables and tables["sales"]["net_sales"].dtype.kind == "f"
    assert tables["sales"]["date"].min() == pd.Timestamp(s["date"].min())       # day-first dates read correctly
    assert "products table built from the sales file" in ingest.check_report(tables, maps)


def test_database_queries_must_be_read_only():
    with pytest.raises(PermissionError):
        ingest.load_database("sqlite://", {"sales": "DELETE FROM sales"})
    with pytest.raises(PermissionError):
        ingest.load_database("sqlite://", {"sales": "SELECT * FROM sales; DROP TABLE sales"})


def test_budget_numbers_add_up(ctx):
    from shelly.reports.retail import budget_specs
    sp = budget_specs(ctx, "All", 3)
    k = {x["label"]: x["value"] for x in sp["kpis"]}
    tbl = next(b for s in sp["sections"] for b in s["blocks"] if b["type"] == "table")
    assert len(tbl["rows"]) == 3
    ly = sum(r["ly"] for r in tbl["rows"])
    assert k["Budget"] == pytest.approx(ly * (1 + ctx.growth), rel=1e-6)
    assert k["Gap to budget"] == pytest.approx(k["Forecast"] - k["Budget"], rel=1e-6)
    assert sp["approvals"], "budgets must ask for approval"


def test_every_report_type_builds(ctx, tmp_path):
    from shelly.reports.build import build_all
    r = build_all(ctx, tmp_path, xlsx=True, with_cache=False, log=lambda *a: None)
    assert len(r["files"]) == 11
    types = {s["type"] for s in r["specs"].values()}
    assert {"budget", "forecast", "weekly", "category", "labour", "stock", "waste", "pack"} <= types
    for f in r["files"].values():
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(f)                      # pivots load (cache ids consistent)
        assert {"Dashboard", "Model", "Pivot", "Raw", "Lookup", "Assumptions", "Notes"} <= set(wb.sheetnames)
        assert len(wb["Pivot"]._pivots) >= 1
        assert any(isinstance(c.value, str) and c.value.startswith("=SUMIFS") for row in wb["Model"].iter_rows() for c in row) or \
            any(isinstance(c.value, str) and c.value.startswith("=") for row in wb["Model"].iter_rows() for c in row)
        assert "SelCat" in wb.defined_names


def test_stock_formulas_block_recalled_items(ctx, tmp_path):
    from shelly.reports.excel import build_workbook
    from shelly.reports.retail import stock_workbook
    plan = stock_workbook(ctx)
    assert plan.raw["Blocked"].sum() >= 1
    p = build_workbook(plan, tmp_path / "s.xlsx")
    wb = openpyxl.load_workbook(p)
    f = [c.value for row in wb["Model"].iter_rows() for c in row if isinstance(c.value, str) and "Blocked" in c.value]
    assert f, "order formula must check the recall block"
