"""Real Excel PivotTables written with openpyxl's pivot model (no Excel needed).

The cache holds every source record, so the pivot works in Excel straight away and refreshes on open;
the summarised values are also written into the sheet cells, so phone and web viewers that don't
run pivots still show the numbers.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import (DataField, FieldItem, Location, PageField, PivotField, PivotTableStyle,
                                  RowColField, RowColItem, TableDefinition)
from openpyxl.pivot.fields import Number, Text, Index
from openpyxl.utils import get_column_letter

_CACHE_ID = [0]


def _is_num(s: pd.Series) -> bool:
    return pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s)


def make_cache(raw: pd.DataFrame, raw_sheet: str, raw_ref: str, cache_id: int = 1) -> dict:
    """One pivot cache per source table; every PivotTable on that table shares it (as Excel does)."""
    df = raw.copy()
    for c in df.columns:                      # dates go into the cache as ISO text
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].dt.strftime("%Y-%m-%d")
    fields = list(df.columns)
    cat_fields = [f for f in fields if not _is_num(df[f])]
    shared: dict[str, list] = {f: sorted(df[f].astype(str).unique().tolist()) for f in cat_fields}

    cfields = []
    for f in fields:
        if f in shared:
            cfields.append(CacheField(name=f, numFmtId=0, sharedItems=SharedItems(
                _fields=[Text(v=v) for v in shared[f]], count=len(shared[f]))))
        else:
            col = pd.to_numeric(df[f], errors="coerce").fillna(0)
            is_int = bool((col == col.round()).all())
            cfields.append(CacheField(name=f, numFmtId=0, sharedItems=SharedItems(
                containsSemiMixedTypes=False, containsString=False, containsNumber=True, containsInteger=is_int or None,
                minValue=float(col.min()), maxValue=float(col.max()))))
    idx = {f: {v: i for i, v in enumerate(shared[f])} for f in shared}
    recs = []
    for row in df.itertuples(index=False):
        items = []
        for f, v in zip(fields, row):
            if f in shared:
                items.append(Index(v=idx[f][str(v)]))
            else:
                items.append(Number(v=float(0 if pd.isna(v) else v)))
        recs.append(Record(_fields=items))
    cache = CacheDefinition(refreshOnLoad=True, recordCount=len(recs), createdVersion=6, refreshedVersion=6,
                            minRefreshableVersion=3,
                            cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref=raw_ref, sheet=raw_sheet)),
                            cacheFields=cfields)
    cache.records = RecordList(r=recs, count=len(recs))
    return {"cache": cache, "df": df, "fields": fields, "shared": shared, "idx": idx, "id": cache_id}


def add_pivot(ws, bundle: dict, anchor_row: int, rows: str, cols: str | None, value: str, page: str | None = None,
              name: str = "PivotTable1", num_fmt: int = 3, value_caption: str | None = None,
              style: str = "PivotStyleMedium2") -> tuple[int, int]:
    """Add a pivot (sum of `value` by `rows` x `cols`, optional page filter) at column A of `ws`.
    Returns (last_row, last_col) of the rendered area."""
    df, fields, shared, idx, cache = bundle["df"], bundle["fields"], bundle["shared"], bundle["idx"], bundle["cache"]

    # ---- rendered summary (what Excel would show), compact layout
    agg = df.pivot_table(index=rows, columns=cols, values=value, aggfunc="sum", fill_value=0) if cols else \
        df.groupby(rows)[value].sum().to_frame(value)
    r_items = [str(v) for v in agg.index]
    c_items = [str(v) for v in agg.columns] if cols else []
    top = anchor_row + (2 if page else 0)
    if page:
        ws.cell(anchor_row, 1, page); ws.cell(anchor_row, 2, "(All)")
    cap = value_caption or f"Sum of {value}"
    ncols = (len(c_items) + 1) if cols else 1
    if cols:
        ws.cell(top, 1, cap); ws.cell(top, 2, "Column Labels")
        ws.cell(top + 1, 1, "Row Labels")
        for j, c in enumerate(c_items):
            ws.cell(top + 1, 2 + j, c)
        ws.cell(top + 1, 2 + len(c_items), "Grand Total")
        for i, r in enumerate(r_items):
            ws.cell(top + 2 + i, 1, r)
            vals = agg.iloc[i].values
            for j, v in enumerate(vals):
                ws.cell(top + 2 + i, 2 + j, float(v))
            ws.cell(top + 2 + i, 2 + len(c_items), float(vals.sum()))
        ws.cell(top + 2 + len(r_items), 1, "Grand Total")
        for j in range(len(c_items)):
            ws.cell(top + 2 + len(r_items), 2 + j, float(agg.iloc[:, j].sum()))
        ws.cell(top + 2 + len(r_items), 2 + len(c_items), float(agg.values.sum()))
        last_row, last_col = top + 2 + len(r_items), 2 + len(c_items)
        loc = Location(ref=f"A{top}:{get_column_letter(last_col)}{last_row}", firstHeaderRow=1, firstDataRow=2,
                       firstDataCol=1, rowPageCount=1 if page else None, colPageCount=1 if page else None)
    else:
        ws.cell(top, 1, "Row Labels"); ws.cell(top, 2, cap)
        for i, r in enumerate(r_items):
            ws.cell(top + 1 + i, 1, r); ws.cell(top + 1 + i, 2, float(agg.iloc[i, 0]))
        ws.cell(top + 1 + len(r_items), 1, "Grand Total"); ws.cell(top + 1 + len(r_items), 2, float(agg.values.sum()))
        last_row, last_col = top + 1 + len(r_items), 2
        loc = Location(ref=f"A{top}:B{last_row}", firstHeaderRow=1, firstDataRow=1, firstDataCol=1,
                       rowPageCount=1 if page else None, colPageCount=1 if page else None)
    for rr in range(top + (2 if cols else 1), last_row + 1):
        for cc in range(2, last_col + 1):
            ws.cell(rr, cc).number_format = "#,##0"

    pfields = []
    for f in fields:
        if f == rows or f == cols or f == page:
            axis = {rows: "axisRow", cols: "axisCol", page: "axisPage"}[f]
            items = [FieldItem(x=i) for i in range(len(shared[f]))] + [FieldItem(t="default")]
            pfields.append(PivotField(axis=axis, showAll=False, items=items))
        elif f == value:
            pfields.append(PivotField(dataField=True, showAll=False))
        else:
            pfields.append(PivotField(showAll=False))

    def items_for(f, present):
        out = [RowColItem(x=[Index(v=idx[f][v])]) for v in present]
        out.append(RowColItem(t="grand", x=[Index(v=0)]))
        return out

    t = TableDefinition(name=name, cacheId=bundle["id"], dataCaption="Values", updatedVersion=6, minRefreshableVersion=3,
                        createdVersion=6, useAutoFormatting=True, itemPrintTitles=True, indent=0, outline=True,
                        outlineData=True, multipleFieldFilters=False, applyNumberFormats=False, applyBorderFormats=False,
                        applyFontFormats=False, applyPatternFormats=False, applyAlignmentFormats=False,
                        applyWidthHeightFormats=True, location=loc, pivotFields=pfields,
                        rowFields=[RowColField(x=fields.index(rows))], rowItems=items_for(rows, r_items),
                        colFields=[RowColField(x=fields.index(cols))] if cols else [],
                        colItems=items_for(cols, c_items) if cols else [RowColItem()],
                        pageFields=[PageField(fld=fields.index(page), hier=-1)] if page else [],
                        dataFields=[DataField(name=cap, fld=fields.index(value), baseField=0, baseItem=0, numFmtId=num_fmt)],
                        pivotTableStyleInfo=PivotTableStyle(name=style, showRowHeaders=True, showColHeaders=True,
                                                            showRowStripes=False, showColStripes=False, showLastColumn=True))
    t.cache = cache
    ws.add_pivot(t)
    return last_row, last_col
