"""
CRISP-DM Phase 6 (Deployment) - hand the results to the tools the business already uses.

  SQLite warehouse + SQL KPI queries   -> outputs/warehouse.db, sql/*.sql
  Power BI star schema + DAX measures  -> outputs/powerbi/
  Tableau tidy extract                 -> outputs/tableau/
  Excel workbook (tables, formulas,
    pivot-ready data, scenario model)  -> outputs/Weekly_Sales_Digest.xlsx
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
SQL_DIR = ROOT / "sql"


# ------------------------------------------------------------------ SQL
def build_warehouse(P, out_dir: Path) -> tuple[Path, dict[str, pd.DataFrame]]:
    db = out_dir / "warehouse.db"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    s = P.sales.copy()
    s["date"] = s["date"].dt.strftime("%Y-%m-%d")
    s[["date", "sku", "channel", "units", "net_sales", "cost", "promo_flag"]].to_sql("fact_sales", con, index=False)
    P.products.to_sql("dim_product", con, index=False)
    if P.waste is not None:
        w = P.waste[["date", "sku", "units_wasted", "units_marked_down", "waste_value"]].copy()
        w["date"] = w["date"].dt.strftime("%Y-%m-%d")
        w.to_sql("fact_waste", con, index=False)
    if P.inventory is not None:
        i = P.inventory[["count_date", "sku", "sap_qty", "physical_qty", "variance_units", "variance_value"]].copy()
        i["count_date"] = i["count_date"].dt.strftime("%Y-%m-%d")
        i.to_sql("fact_stock_count", con, index=False)

    results = {}
    asof = P.asof.strftime("%Y-%m-%d")
    for f in sorted(SQL_DIR.glob("*.sql")):
        q = f.read_text().replace(":asof", f"'{asof}'")
        try:
            results[f.stem] = pd.read_sql_query(q, con)
        except Exception as e:  # keep the run going, report the failure
            results[f.stem] = pd.DataFrame({"error": [str(e)]})
    con.close()
    return db, results


# ------------------------------------------------------------------ Power BI
DAX = """// ---- Paste into Power BI: Modeling > New measure ----
// Model: fact_sales[sku] -> dim_product[sku], fact_sales[date] -> dim_date[date],
//        fact_sales[channel] -> dim_channel[channel]; mark dim_date as date table.

Total Sales = SUM ( fact_sales[net_sales] )
Total COGS = SUM ( fact_sales[cost] )
Gross Margin = [Total Sales] - [Total COGS]
GM % = DIVIDE ( [Gross Margin], [Total Sales] )
Units = SUM ( fact_sales[units] )
Avg Selling Price = DIVIDE ( [Total Sales], [Units] )

Sales PW = CALCULATE ( [Total Sales], DATEADD ( dim_date[date], -7, DAY ) )
WoW % = DIVIDE ( [Total Sales] - [Sales PW], [Sales PW] )
Sales LY = CALCULATE ( [Total Sales], DATEADD ( dim_date[date], -364, DAY ) )
YoY % = DIVIDE ( [Total Sales] - [Sales LY], [Sales LY] )
Budget = [Sales LY] * ( 1 + 0.04 )
Var to Budget = [Total Sales] - [Budget]

Delivery Sales = CALCULATE ( [Total Sales], dim_channel[is_delivery] = TRUE () )
Delivery Share % = DIVIDE ( [Delivery Sales], [Total Sales] )
Promo Share % = DIVIDE ( CALCULATE ( [Total Sales], fact_sales[promo_flag] = 1 ), [Total Sales] )

Waste Value = SUM ( fact_waste[waste_value] )
Waste % Sales = DIVIDE ( [Waste Value], [Total Sales] )
Shrink Value = SUM ( fact_stock_count[variance_value] )

Sales 7D Rolling = CALCULATE ( [Total Sales], DATESINPERIOD ( dim_date[date], MAX ( dim_date[date] ), -7, DAY ) )
Category Rank = RANKX ( ALL ( dim_product[category] ), [Total Sales] )
ABC Class =
VAR cum = SUMX ( FILTER ( ALL ( dim_product[sku] ), [Total Sales] >= CALCULATE ( [Total Sales] ) ), [Total Sales] )
VAR share = DIVIDE ( cum, CALCULATE ( [Total Sales], ALL ( dim_product ) ) )
RETURN SWITCH ( TRUE (), share <= 0.8, "A", share <= 0.95, "B", "C" )
"""


def export_powerbi(P, R, out_dir: Path) -> Path:
    d = out_dir / "powerbi"
    d.mkdir(parents=True, exist_ok=True)
    s = P.sales
    s[["date", "sku", "channel", "units", "net_sales", "cost", "promo_flag"]].to_csv(d / "fact_sales.csv", index=False)
    P.products.to_csv(d / "dim_product.csv", index=False)
    dates = pd.DataFrame({"date": pd.date_range(s["date"].min(), P.asof + pd.Timedelta(days=14))})
    dates["year"], dates["month"] = dates["date"].dt.year, dates["date"].dt.month
    dates["month_name"] = dates["date"].dt.strftime("%b")
    dates["iso_week"] = dates["date"].dt.isocalendar().week.astype(int)
    dates["weekday"] = dates["date"].dt.strftime("%a")
    dates["is_weekend"] = dates["date"].dt.dayofweek >= 5
    dates.to_csv(d / "dim_date.csv", index=False)
    pd.DataFrame({"channel": ["in_store", "uber_eats", "on_demand"],
                  "channel_label": ["In-store", "Uber Eats", "On-Demand"],
                  "is_delivery": [False, True, True]}).to_csv(d / "dim_channel.csv", index=False)
    if P.waste is not None:
        P.waste[["date", "sku", "units_wasted", "units_marked_down", "waste_value"]].to_csv(d / "fact_waste.csv", index=False)
    if P.inventory is not None:
        P.inventory[["count_date", "sku", "sap_qty", "physical_qty", "variance_units", "variance_value"]].to_csv(
            d / "fact_stock_count.csv", index=False)
    R["forecast"]["forecast"].to_csv(d / "fact_forecast.csv", index=False)
    R["segments"]["sku_segments"][["sku", "segment", "pc1", "pc2"]].to_csv(d / "dim_segment.csv", index=False)
    (d / "measures.dax").write_text(DAX)
    return d


# ------------------------------------------------------------------ Tableau
def export_tableau(P, R, out_dir: Path) -> Path:
    d = out_dir / "tableau"
    d.mkdir(parents=True, exist_ok=True)
    t = P.sales.merge(R["segments"]["sku_segments"][["sku", "segment"]], on="sku", how="left")
    t["gross_margin"] = t["net_sales"] - t["cost"]
    t["is_delivery"] = t["channel"].isin(["uber_eats", "on_demand"])
    t["weekday"] = t["date"].dt.strftime("%a")
    t[["date", "weekday", "sku", "product_name", "category", "segment", "supplier", "channel",
       "is_delivery", "promo_flag", "units", "net_sales", "cost", "gross_margin"]].to_csv(
        d / "sales_extract.csv", index=False)
    R["forecast"]["forecast"].to_csv(d / "forecast_extract.csv", index=False)
    return d


# ------------------------------------------------------------------ Excel
HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
WHITE_B = Font(color="FFFFFF", bold=True)


def _sheet_from_df(wb, title, df, table_name=None, money_cols=(), pct_cols=()):
    ws = wb.create_sheet(title)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for c in ws[1]:
        c.fill, c.font = HEAD, WHITE_B
    for i, col in enumerate(df.columns, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = min(max(len(str(col)) + 2, 12), 40)
        fmt = "$#,##0" if col in money_cols else ("0.0" if col in pct_cols else None)
        if fmt:
            for cell in ws[L][1:]:
                cell.number_format = fmt
    if table_name and len(df):
        ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    ws.freeze_panes = "A2"
    return ws


def export_excel(P, R, sql_results: dict, out_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    k = R["commercial"]["kpis"]

    # 1. KPI summary
    ws = wb.create_sheet("Summary")
    ws["A1"], ws["A1"].font = f"Shelly v0.1 · Weekly Sales Digest - week ending {P.asof:%d %b %Y}", Font(bold=True, size=14)
    rows = [("Metric", "This week", "Last week", "Same week LY"),
            ("Sales", k["current"]["sales"], k["previous"]["sales"], k["last_year"]["sales"]),
            ("Units", k["current"]["units"], k["previous"]["units"], k["last_year"]["units"]),
            ("Gross margin", k["current"]["gross_margin"], k["previous"]["gross_margin"], k["last_year"]["gross_margin"]),
            ("Waste + markdown", k["current"]["waste_value"], k["previous"]["waste_value"], None),
            ("Budget", k["budget"], None, None),
            ("GM %", "=B6/B4", "=C6/C4", "=D6/D4"),
            ("WoW %", "=B4/C4-1", None, None),
            ("YoY %", "=B4/D4-1", None, None),
            ("vs Budget %", "=B4/B8-1", None, None)]
    for r_i, r in enumerate(rows, 3):          # header on row 3, data rows 4-12
        for c_i, v in enumerate(r, 1):
            c = ws.cell(r_i, c_i, None if v is None else (float(v) if hasattr(v, "item") else v))
            if r_i == 3:
                c.fill, c.font = HEAD, WHITE_B
            elif c_i > 1:
                c.number_format = "0.0%" if r_i >= 9 else ("#,##0" if r_i == 5 else "$#,##0")
    ws.column_dimensions["A"].width = 22
    for L in "BCD":
        ws.column_dimensions[L].width = 16

    # 2. Category P&L with live formulas
    pl = R["commercial"]["category_pl"].reset_index()[
        ["category", "sales", "cogs", "waste_value", "markdown_value", "sales_prev", "sales_ly", "budget"]]
    ws = _sheet_from_df(wb, "Category P&L", pl, None, money_cols=pl.columns[1:])
    extra = ["Gross margin", "GM %", "Contribution", "WoW %", "YoY %", "Var to budget"]
    base = len(pl.columns)
    for j, h in enumerate(extra, 1):
        c = ws.cell(1, base + j, h)
        c.fill, c.font = HEAD, WHITE_B
        ws.column_dimensions[get_column_letter(base + j)].width = 14
    for i in range(2, len(pl) + 2):
        f = {"Gross margin": (f"=B{i}-C{i}", "$#,##0"), "GM %": (f"=IFERROR(I{i}/B{i},0)", "0.0%"),
             "Contribution": (f"=I{i}-D{i}-E{i}", "$#,##0"), "WoW %": (f"=IFERROR(B{i}/F{i}-1,0)", "0.0%"),
             "YoY %": (f"=IFERROR(B{i}/G{i}-1,0)", "0.0%"), "Var to budget": (f"=B{i}-H{i}", "$#,##0")}
        for j, h in enumerate(extra, 1):
            c = ws.cell(i, base + j, f[h][0])
            c.number_format = f[h][1]
    t = len(pl) + 2
    ws.cell(t, 1, "TOTAL").font = Font(bold=True)
    for col in range(2, base + 1):
        L = get_column_letter(col)
        ws.cell(t, col, f"=SUM({L}2:{L}{t - 1})").number_format = "$#,##0"
    ws.cell(t, base + 1, f"=B{t}-C{t}").number_format = "$#,##0"
    ws.cell(t, base + 2, f"=I{t}/B{t}").number_format = "0.0%"

    # 3. Scenario model (what-if) - inputs in yellow drive formulas
    ws = wb.create_sheet("Scenario Model")
    ws["A1"], ws["A1"].font = "What-if scenario: price, volume and waste levers (edit yellow cells)", Font(bold=True, size=12)
    inputs = [("Weekly sales (base)", k["current"]["sales"]), ("Gross margin % (base)", k["current"]["gm_pct"] / 100),
              ("Waste + markdown (base)", k["current"]["waste_value"]), ("Price change %", 0.02),
              ("Price elasticity of demand", -1.2), ("Waste reduction %", 0.25),
              ("Delivery outage days avoided / wk", 1), ("Avg delivery sales per day", k["current"]["sales"] * k["current"]["delivery_share_pct"] / 100 / 7)]
    for i, (lab, v) in enumerate(inputs, 3):
        ws.cell(i, 1, lab)
        c = ws.cell(i, 2, v)
        c.fill = INPUT
        c.number_format = ("0.0%" if "%" in lab else "0.00" if "elasticity" in lab
                           else "0" if "days" in lab else "$#,##0")
    outs = [("Volume change %", "=B6*B7", "0.0%"),
            ("New weekly sales", "=B3*(1+B6)*(1+B12)", "$#,##0"),
            ("New gross margin %", "=1-(1-B4)/(1+B6)", "0.0%"),
            ("New gross margin $", "=B13*B14", "$#,##0"),
            ("Base gross margin $", "=B3*B4", "$#,##0"),
            ("Waste saving $", "=B5*B8", "$#,##0"),
            ("Recovered delivery sales $", "=B9*B10", "$#,##0"),
            ("Weekly GM uplift $", "=B15-B16+B17+B18*B4", "$#,##0"),
            ("Annualised GM uplift $", "=B19*52", "$#,##0")]
    for i, (lab, fml, fmt) in enumerate(outs, 12):
        ws.cell(i, 1, lab).font = Font(bold=True)
        c = ws.cell(i, 2, fml)
        c.number_format = fmt
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 16

    # 4+. data tabs (Excel tables -> PivotTable-ready)
    fc = R["forecast"]["forecast"].copy()
    fc["date"] = fc["date"].dt.date
    _sheet_from_df(wb, "Forecast", fc, "tblForecast", money_cols=["forecast_sales"], pct_cols=["backtest_wape"])
    _sheet_from_df(wb, "Model Leaderboard", R["forecast"]["model_leaderboard"].round(2), "tblModels")
    an = R["anomalies"]
    if len(an):
        _sheet_from_df(wb, "Exceptions", an, "tblExceptions", money_cols=["impact_nzd"])
    rp = R["replenishment"][["sku", "product_name", "category", "supplier", "mu", "safety_stock", "reorder_point",
                             "est_on_hand", "days_cover", "suggested_order", "order_value", "status"]].round(1)
    _sheet_from_df(wb, "Reorder List", rp, "tblReorder", money_cols=["order_value"])
    lp = R["labour"].copy()
    lp["date"] = lp["date"].dt.date
    _sheet_from_df(wb, "Roster Plan", lp.round(1), "tblRoster", money_cols=["forecast_sales", "wage_cost"])
    seg = R["segments"]["sku_segments"][["sku", "product_name", "category", "segment", "avg_daily_units", "gm_pct",
                                         "demand_cv", "waste_rate_pct", "delivery_share_pct", "promo_uplift"]].round(2)
    _sheet_from_df(wb, "Segments", seg, "tblSegments")
    _sheet_from_df(wb, "Range Review", R["assortment"].round(1), "tblRange", money_cols=["sales", "cost", "gm"])
    pivot = P.sales[P.sales["date"] > P.asof - pd.Timedelta(days=91)][
        ["date", "sku", "product_name", "category", "channel", "promo_flag", "units", "net_sales", "cost"]].copy()
    pivot["date"] = pivot["date"].dt.date
    _sheet_from_df(wb, "Pivot Data (13 wks)", pivot, "tblSales13w", money_cols=["net_sales", "cost"])
    for name, df in sql_results.items():
        _sheet_from_df(wb, f"SQL {name}"[:31], df.round(2))
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="center")
    wb.save(out_path)
    return out_path
